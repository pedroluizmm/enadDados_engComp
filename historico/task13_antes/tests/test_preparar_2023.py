"""Regressões da auditoria: domínio após linhas vazias e nota zero versus ausência."""
import io
from pathlib import Path
import sys
import unittest
import zipfile

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from preparar_2023 import dictionary, valid_scores


class AuditRegressionTests(unittest.TestCase):
    def test_domain_after_blank_rows_and_reordered_sheets(self):
        book = openpyxl.Workbook()
        book.active.title = 'Outra aba'
        sheet = book.create_sheet('Dicionário de Variáveis')
        sheet.append(['NOME', 'TIPO', 'TAMANHO', 'DESCRIÇÃO', 'CATEGORIAS'])
        sheet.append(['CO_CATEGAD', 'N', 1, 'Categoria', '4. Privada com fins lucrativos'])
        sheet['E3'] = '5. Privada sem fins lucrativos'
        sheet['E22'] = '8. Comunitária/Confessional'
        sheet['A22'] = ''
        sheet['A23'] = 'CO_ORGACAD'
        sheet['E23'] = '10028 = Universidade'
        workbook, archive = io.BytesIO(), io.BytesIO()
        book.save(workbook)
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('dicionario.xlsx', workbook.getvalue())
        with zipfile.ZipFile(io.BytesIO(archive.getvalue())) as z:
            result = dictionary(z, enade=True)['abas']['Dicionário de Variáveis']
        self.assertEqual(result['CO_CATEGAD']['dominio']['8'], 'Comunitária/Confessional')
        self.assertNotIn('10028', result['CO_CATEGAD']['dominio'])
        self.assertIn(22, result['CO_CATEGAD']['linhas'])

    def test_letter_domains_after_empty_string(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Dicionário de Variáveis'
        sheet.append(['NOME', 'TIPO', 'TAMANHO', 'DESCRIÇÃO', 'CATEGORIAS'])
        sheet.append(['QE_I10', 'C', 1, 'Trabalho', 'A = Não trabalho'])
        sheet.append(['', None, None, None, 'B = Eventualmente'])
        sheet.append(['', None, None, None, 'E = 40 horas ou mais'])
        workbook, archive = io.BytesIO(), io.BytesIO()
        book.save(workbook)
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('dicionario.xlsx', workbook.getvalue())
        with zipfile.ZipFile(io.BytesIO(archive.getvalue())) as z:
            entry = dictionary(z, True)['abas']['Dicionário de Variáveis']['QE_I10']
        self.assertEqual(len(entry['texto_dominio']), 3)

    def test_zero_is_valid_blank_and_absence_are_not(self):
        data = pd.DataFrame({'TP_PRES': ['555', '222', '555', '888'],
                             'TP_PR_GER': ['333', '222', '555', '888'],
                             'NT_GER': [0, None, None, 40],
                             'NT_FG': [0, None, None, 40],
                             'NT_CE': [0, None, None, 40]})
        self.assertEqual(valid_scores(data).tolist(), [True, False, False, False])


if __name__ == '__main__':
    unittest.main()
